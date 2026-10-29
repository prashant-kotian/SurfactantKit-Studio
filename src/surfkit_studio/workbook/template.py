"""The blank template a researcher downloads, fills in and uploads."""
from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import layout as L

BOLD = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")     # cells the researcher types into
CALC_FILL = PatternFill("solid", fgColor="EAEAEA")      # cells Excel calculates

README = [
    "SurfactantKit Studio: binary surfactant workbook",
    "",
    "1. Info: names, surfactant types, temperature, added salt, solvent surface tension.",
    "2. Surfactant 1 and Surfactant 2: the pure surfactants. Surface tension in up to three runs per "
    "concentration (mean and SD are calculated for you); conductivity optional (ionic surfactants).",
    "3. Mix 1 to Mix 5: one sheet per composition. Type alpha1, the mole fraction of surfactant 1 in the "
    "total surfactant, in cell B2. Copy a Mix sheet (and name it Mix 6, Mix 7, ...) for more compositions.",
    "4. Temperature series (optional): CMC of each pure surfactant at 3 or more temperatures.",
    "",
    "Concentrations in mM. Surface tension in mN/m. Yellow cells are yours, grey cells are calculated.",
    "Leave unused rows empty. Do not change the sheet layout.",
    "",
    "Upload the filled file to the Studio (Workbook mode). You get back a completed workbook with the "
    "toolkit's results, Excel Solver sheets set up the usual way (Rubingh and Rosen), a comparison of the "
    "two, data flags, and every plot as PNG and as an xmgrace .agr file.",
]


def _widths(ws, widths: dict[int, float]):
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def add_info_sheet(wb, values: dict | None = None):
    ws = wb.create_sheet(L.INFO)
    ws["A1"] = "Experiment information"; ws["A1"].font = Font(bold=True, size=12)
    for key, (row, label, default, allowed) in L.INFO_FIELDS.items():
        ws.cell(row, 1, label).font = BOLD
        c = ws.cell(row, L.INFO_VALUE_COL, (values or {}).get(key, default))
        c.fill = INPUT_FILL
        if allowed:
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=False)
            ws.add_data_validation(dv); dv.add(c)
    ws["D4"] = ("ionic: Gibbs n = 2 without salt, n = 1 with excess salt; ionic divalent (gemini or dianion, two "
                "counterions): n = 3 without salt; nonionic: n = 1.")
    ws["D11"] = "Rosen beta-sigma is read where the surface tension is (solvent - this value); 20 is the usual choice."
    ws["D13"] = "Display only: every value is stored and calculated at full precision."
    ws["D14"] = ("Used for every CMC, beta and Solver sheet in the output; both the regression-intersection and the "
                 "Szyszkowski CMC are always listed on the Results sheets.")
    _widths(ws, {1: 44, 2: 18, 4: 90})
    return ws


def add_data_sheet(wb, name: str, title: str, alpha: float | None = None):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=12)
    if alpha is not None:
        ws["A2"] = "alpha1"; ws["A2"].font = BOLD
        ws[L.ALPHA_CELL] = alpha; ws[L.ALPHA_CELL].fill = INPUT_FILL
        ws["C2"] = "mole fraction of surfactant 1 in the total surfactant"
    ws.cell(3, L.T_COLS["conc"], "Surface tension").font = BOLD
    ws.cell(3, L.K_COLS["conc"], "Conductivity (optional)").font = BOLD
    for j, h in enumerate(L.T_HEADERS):
        c = ws.cell(L.HEADER_ROW, L.T_COLS["conc"] + j, h); c.font = BOLD
        c.alignment = Alignment(wrap_text=True)
    for j, h in enumerate(L.K_HEADERS):
        c = ws.cell(L.HEADER_ROW, L.K_COLS["conc"] + j, h); c.font = BOLD
        c.alignment = Alignment(wrap_text=True)
    for r in range(L.FIRST_ROW, L.LAST_ROW + 1):
        a, cr, er = f"A{r}", f"C{r}", f"E{r}"
        ws.cell(r, L.T_COLS["log"], f'=IF({a}="","",LOG10({a}))').fill = CALC_FILL
        ws.cell(r, L.T_COLS["mean"], f'=IF(COUNT({cr}:{er})=0,"",AVERAGE({cr}:{er}))').fill = CALC_FILL
        ws.cell(r, L.T_COLS["sd"], f'=IF(COUNT({cr}:{er})<2,"",STDEV({cr}:{er}))').fill = CALC_FILL
        for key in ("conc", "run1", "run2", "run3"):
            ws.cell(r, L.T_COLS[key]).fill = INPUT_FILL
        for key in ("conc", "kappa"):
            ws.cell(r, L.K_COLS[key]).fill = INPUT_FILL
    ws.row_dimensions[L.HEADER_ROW].height = 45
    _widths(ws, {1: 11, 2: 9, 3: 13, 4: 11, 5: 11, 6: 11, 7: 10, 9: 11, 10: 14})
    ws.freeze_panes = ws.cell(L.FIRST_ROW, 1)
    return ws


def add_temperature_sheet(wb):
    ws = wb.create_sheet(L.TEMPERATURE)
    ws["A1"] = "Temperature series (optional): CMC of each pure surfactant at 3 or more temperatures"
    ws["A1"].font = Font(bold=True, size=12)
    for j, h in enumerate(["Temperature (C)", "CMC surfactant 1 (mM)", "CMC surfactant 2 (mM)"]):
        ws.cell(L.HEADER_ROW, 1 + j, h).font = BOLD
    for r in range(L.TEMP_FIRST_ROW, L.TEMP_FIRST_ROW + L.TEMP_N_ROWS):
        for col in (1, 2, 3):
            ws.cell(r, col).fill = INPUT_FILL
    _widths(ws, {1: 16, 2: 22, 3: 22})
    return ws


def new_workbook(info: dict | None = None, mixes: list[tuple[str, float]] | None = None):
    """mixes: (sheet name, alpha1) pairs; default Mix 1..5 at alpha1 0.1..0.9."""
    if mixes is None:
        mixes = list(zip(L.mix_sheet_names(), L.DEFAULT_ALPHAS))
    wb = openpyxl.Workbook()
    rd = wb.active; rd.title = "Read me"
    for i, line in enumerate(README, 1):
        rd.cell(i, 1, line)
    rd["A1"].font = Font(bold=True, size=12); rd.column_dimensions["A"].width = 120
    add_info_sheet(wb, info)
    for i, name in enumerate(L.PURE):
        add_data_sheet(wb, name, f"Pure {name.lower()}: name and type on Info")
    for name, a in mixes:
        add_data_sheet(wb, name, "Mixture of surfactant 1 and 2", alpha=a)
    add_temperature_sheet(wb)
    return wb


def build_template() -> bytes:
    buf = io.BytesIO()
    new_workbook().save(buf)
    return buf.getvalue()
