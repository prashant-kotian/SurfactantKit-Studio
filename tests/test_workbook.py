"""Workbook mode: template, reader, analysis, completed workbook, plots."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from surfactantkit import curve_analysis as ca

import synthetic_workbook as S
from surfkit_studio.workbook import WorkbookError, analyse, build_output_zip, build_template, read_workbook
from surfkit_studio.workbook import layout as L
from surfkit_studio.workbook.analysis import global_beta_least_squares
from surfkit_studio.workbook.figures import build_figs, render_agr
from surfkit_studio.workbook.template import new_workbook
from surfkit_studio.workbook.writer import sig_format

EX = Path(__file__).resolve().parents[1] / "examples"


def _book(series: dict, info=None, mixes=None) -> bytes:
    """series: sheet -> (conc, [run lists] or gammas, optional conductivity pairs)."""
    wb = new_workbook(info, mixes)
    for sheet, data in series.items():
        conc, gam = data[0], data[1]
        for i, (c, g) in enumerate(zip(conc, gam)):
            wb[sheet].cell(L.FIRST_ROW + i, L.T_COLS["conc"], c)
            for j, x in enumerate(g if isinstance(g, (list, tuple)) else [g]):
                wb[sheet].cell(L.FIRST_ROW + i, L.T_COLS["run1"] + j, x)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ template
def test_template_has_every_sheet_and_live_formulas():
    wb = openpyxl.load_workbook(io.BytesIO(build_template()))
    assert wb.sheetnames == ["Read me", "Info", "Surfactant 1", "Surfactant 2", "Mix 1", "Mix 2", "Mix 3",
                             "Mix 4", "Mix 5", "Temperature series"]
    ws = wb["Mix 3"]
    assert ws[L.ALPHA_CELL].value == 0.5
    r = L.FIRST_ROW
    assert ws.cell(r, L.T_COLS["mean"]).value == f'=IF(COUNT(C{r}:E{r})=0,"",AVERAGE(C{r}:E{r}))'
    assert ws.cell(r, L.T_COLS["sd"]).value == f'=IF(COUNT(C{r}:E{r})<2,"",STDEV(C{r}:E{r}))'
    assert ws.cell(r, L.T_COLS["log"]).value == f'=IF(A{r}="","",LOG10(A{r}))'


def test_blank_template_reads_back_with_defaults_and_no_data():
    st = read_workbook(build_template())
    assert st.info["gamma_solvent"] == 72.0 and st.info["sig_figs"] == 2 and st.info["s1_type"] == "ionic"
    assert all(d.tension is None for d in st.pure) and st.mixes == [] and st.temperature == []


def test_triplicate_mean_and_sd_are_recomputed_not_trusted():
    data = _book({"Surfactant 1": ([0.1, 0.2], [[70.0, 70.2, 70.4], [65.0, 65.0]])})
    t = read_workbook(data).pure[0].tension
    assert t.mean == pytest.approx([70.2, 65.0]) and t.sd[0] == pytest.approx(0.2) and t.sd[1] == 0.0
    one = read_workbook(_book({"Surfactant 1": ([0.1], [[70.0]])}))
    assert one.pure[0].tension.sd == [None] and any("single run" in n for n in one.notes)


@pytest.mark.parametrize("series,msg", [
    ({"Surfactant 1": ([0.1, 0.1], [70.0, 69.0])}, "same concentration"),
    ({"Surfactant 1": ([0.1], [0.07])}, "10-80 mN/m"),
    ({"Surfactant 1": ([-0.1], [70.0])}, "must be positive"),
])
def test_reader_rejects_bad_tables_naming_the_cell(series, msg):
    with pytest.raises(WorkbookError, match=msg):
        read_workbook(_book(series))


def test_reader_rejects_a_file_that_is_not_the_template():
    wb = openpyxl.Workbook(); buf = io.BytesIO(); wb.save(buf)
    with pytest.raises(WorkbookError, match="not the Studio template"):
        read_workbook(buf.getvalue())
    with pytest.raises(WorkbookError, match="readable"):
        read_workbook(b"not a zip")


# ------------------------------------------------------------------ analysis
def test_synthetic_study_recovers_the_parameters_it_was_built_from():
    res = analyse(read_workbook(S.build()))
    t = S.truth()
    assert res.rubingh["beta_mean"] == pytest.approx(t["beta_m"], abs=0.01)
    assert res.rubingh["beta_global"] == pytest.approx(t["beta_m"], abs=0.01)
    assert res.rubingh["model"] == "rubingh_regular_solution"
    assert res.rosen["beta_mean"] == pytest.approx(t["beta_sigma"], rel=0.03)   # pre-CMC line on a curved isotherm
    p1, p2 = res.pure
    assert p1["cmc_k"] == pytest.approx(t["cmc1"], rel=1e-3) and p2["cmc_k"] == pytest.approx(t["cmc2"], rel=1e-3)
    assert p1["beta_cb"] == pytest.approx(1 - 30 / 80, abs=1e-6)
    assert p1["cmc_t"] == pytest.approx(t["cmc1"], rel=0.08) and p2["cmc_t"] == pytest.approx(t["cmc2"], rel=0.08)
    assert [s["result"] for s in res.synergism] == [True, True, True]
    assert not [f for f in res.flags if f.level == "warning"]


def test_global_beta_reproduces_the_ogino_abe_sds_poe_example():
    # Ogino and Abe, Mixed Surfactant Systems (1992) p. 13: SDS (8 mM) + C18POE20 (0.01 mM).
    # Least squares on the 1/CMC residual gives -7.58; the workbook's own composition-residual table
    # picked -7, which is Solver noise (that residual is ~0 for any beta).
    b, ssd, edge = global_beta_least_squares([0.2, 0.4, 0.6, 0.8], [0.011, 0.012, 0.015, 0.023], 8.0, 0.01)
    assert b == pytest.approx(-7.584, abs=0.01) and not edge


USER_DUMMY_PURE = ([0.056, 0.0839, 0.1257, 0.1883, 0.2821, 0.4226, 0.633, 0.9484, 1.4208, 2.1284, 3.1886, 4.7769,
                    7.1563, 10.7208, 10.8, 16.0609, 24.0608, 36.0456, 54],
                   [71, 70.93, 70.74, 70.69, 70.68, 70.6, 70.58, 70.41, 50.46, 47.92, 45.45, 42.92, 40.33, 37.89,
                    37.81, 37.78, 37.81, 37.81, 37.79])
USER_DUMMY_MIX = ([0.00144, 0.00228, 0.00362, 0.00573, 0.00909, 0.0144, 0.02282, 0.03617, 0.05733, 0.072, 0.09086,
                   0.144, 0.22823, 0.36171, 0.57327, 0.90858, 1.44],
                  [71.54, 71.36, 71.18, 71.18, 71.04, 70.95, 70.9, 70.72, 70.74, 32.02, 31.9, 31.96, 31.84, 31.96,
                   31.91, 31.96, 31.88])


def test_step_shaped_data_is_flagged_and_kept_out_of_beta():
    res = analyse(read_workbook(_book({"Surfactant 1": USER_DUMMY_PURE, "Surfactant 2": USER_DUMMY_PURE,
                                       "Mix 1": USER_DUMMY_MIX})))
    warn = [f for f in res.flags if f.level == "warning"]
    assert any(f.where == "Surfactant 1" and "falls 19.9 mN/m" in f.message for f in warn)
    assert any(f.where == "Mix 1" and ("falls" in f.message or "almost no surface tension decline" in f.message)
               for f in warn)
    assert "beta_mean" not in res.rubingh and "beta_mean" not in res.rosen


def test_workbook_cmc_equals_the_studio_tensiometry_technique_on_real_sds_data():
    rows = list(csv.reader(open(EX / "SDS_surface_tension.csv", encoding="utf-8")))[1:]
    c, g = [float(r[0]) for r in rows], [float(r[1]) for r in rows]
    res = analyse(read_workbook(_book({"Surfactant 1": (c, g)})))
    assert res.pure[0]["cmc_t"] == ca.cmc_from_surface_tension_curve(c, g).cmc_mM


# ------------------------------------------------------------------ output
@pytest.fixture(scope="module")
def output():
    z, summary = build_output_zip(S.build(), "synthetic")
    zf = zipfile.ZipFile(io.BytesIO(z))
    return zf, summary, openpyxl.load_workbook(io.BytesIO(zf.read("synthetic_results.xlsx")))


def test_zip_holds_workbook_and_every_plot_as_png_and_agr(output):
    zf, summary, _ = output
    names = zf.namelist()
    pngs = {n[:-4] for n in names if n.endswith(".png")}
    agrs = {n[:-4] for n in names if n.endswith(".agr")}
    assert pngs == agrs and len(pngs) == summary["n_plots"] == 12
    assert zf.read("plots/all_isotherms.png")[:8] == b"\x89PNG\r\n\x1a\n"


def test_solver_sheets_carry_a_preset_model_with_qualified_ranges(output):
    _, _, wb = output
    for sheet in ("Rubingh per-point (Solver)", "Rubingh global (Solver)", "Rosen per-point (Solver)"):
        names = wb[sheet].defined_names
        q = f"'{sheet}'!"
        assert names["solver_typ"].attr_text == "2" and names["solver_opt"].attr_text.startswith(q)
        assert all(part.startswith(q) for part in names["solver_adj"].attr_text.split(","))
    g = wb["Rubingh global (Solver)"]
    assert g["B6"].value == -1.0 and g["P9"].value == "=(M9-O9)^2" and g["J9"].value == "=A9-I9"
    assert wb["Rubingh global (Solver)"].defined_names["solver_rel1"].attr_text == "2"   # Solver cells = 0


def test_comparison_formulas_point_at_existing_sheets(output):
    _, _, wb = output
    refs = set()
    for row in wb["Comparison"].iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                refs |= set(re.findall(r"'([^']+)'!", c.value))
    assert refs and refs <= set(wb.sheetnames)


def test_results_display_the_chosen_significant_figures(output):
    _, _, wb = output
    ws = wb["Results - pure"]
    cell = ws["B5"]                                    # CMC of surfactant 1, about 7.6 mM
    assert isinstance(cell.value, float) and cell.number_format == "0.0"   # 2 significant figures, full value kept
    # number formats: 37.83 shows "38", 0.0123 shows "0.012", 2.4e-6 shows "2.4E-06"
    assert sig_format(37.83, 2) == "0" and sig_format(0.0123, 2) == "0.000" and sig_format(2.4e-6, 2) == "0.0E+00"
    assert sig_format(-7.58, 2) == "0.0" and sig_format(-7.58, 3) == "0.00"


def test_agr_has_one_set_per_series_and_error_bars_where_runs_exist():
    st = read_workbook(S.build())
    res = analyse(st)
    fig = next(f for f in build_figs(st, res) if f.name == "surfactant_1_tension")
    agr = render_agr(fig)
    assert agr.count("@target G0.S") == len(fig.series)
    assert "@type xydy" in agr and "@    xaxes scale Logarithmic" in agr
    first = agr.split("@type xydy\n")[1].splitlines()[0].split()
    assert len(first) == 3 and float(first[0]) == pytest.approx(st.pure[0].tension.conc[0])


# ------------------------------------------------------------------ web endpoints
def test_endpoints_serve_the_template_and_return_the_results_zip():
    import json

    from fastapi.testclient import TestClient

    from surfkit_studio.app import app
    client = TestClient(app)
    t = client.get("/api/workbook/template.xlsx")
    assert t.status_code == 200 and t.content[:2] == b"PK"
    r = client.post("/api/workbook/analyze", files={"file": ("my study.xlsx", S.build(), "application/octet-stream")})
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    assert 'filename="my_study_results.zip"' in r.headers["content-disposition"]
    s = json.loads(r.headers["x-summary"])
    assert round(s["beta_m_mean"], 2) == -6.0 and s["n_plots"] == 12
    bad = client.post("/api/workbook/analyze", files={"file": ("x.xlsx", b"nope", "application/octet-stream")})
    assert bad.status_code == 400 and "readable" in bad.json()["detail"]
