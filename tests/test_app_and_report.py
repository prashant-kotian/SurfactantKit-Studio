"""End-to-end through the HTTP API, plus report formatting rules."""
import io
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from surfkit_studio.app import app

EX = Path(__file__).resolve().parents[1] / "examples"
client = TestClient(app)


def _session():
    return client.post("/api/session").json()["sid"]


def _upload(sid, tech, name):
    r = client.post("/api/upload", data={"sid": sid, "technique": tech},
                    files={"file": (name, (EX / name).read_bytes(), "text/csv")})
    assert r.status_code == 200, r.text
    return r.json()["mapping"]


@pytest.fixture(scope="module")
def full_session():
    sid = _session()
    m = _upload(sid, "cmc_tensiometry", "SDS_surface_tension.csv")
    r = client.post("/api/run", json={"sid": sid, "technique": "cmc_tensiometry", "method": "rss", "mapping": m,
                                      "params": {"temperature_K": 298.15, "system_type": "ionic_no_added_salt"}})
    assert r.status_code == 200, r.text
    assert {n["id"] for n in r.json()["ready_next"]} >= {"surface_excess", "gibbs_energy_micellization"}
    m = _upload(sid, "cmc_conductivity", "synthetic_conductivity.csv")
    assert client.post("/api/run", json={"sid": sid, "technique": "cmc_conductivity", "method": "rss",
                                         "mapping": m}).status_code == 200
    for tech, method in (("surface_excess", "gibbs"), ("isotherm_selection", "bic"),
                         ("gibbs_energy_micellization", "pseudo_phase_ionic")):
        cat = client.get(f"/api/catalog?sid={sid}").json()["techniques"][tech]
        params = {p["key"]: p["value"] for p in cat["params"]}       # prefilled from the session
        r = client.post("/api/run", json={"sid": sid, "technique": tech, "method": method, "params": params})
        assert r.status_code == 200, (tech, r.text)
    return sid


def test_derived_technique_blocked_until_its_inputs_exist():
    sid = _session()
    r = client.post("/api/run", json={"sid": sid, "technique": "surface_excess", "method": "gibbs",
                                      "params": {"temperature_K": 298.15, "system_type": "nonionic"}})
    assert r.status_code == 400 and "first" in r.json()["detail"]


def test_session_values_prefill_derived_parameters(full_session):
    cat = client.get(f"/api/catalog?sid={full_session}").json()["techniques"]
    beta = next(p for p in cat["gibbs_energy_micellization"]["params"] if p["key"] == "beta")
    assert beta["value"] == pytest.approx(0.70, abs=0.03)        # carried over from conductivity


def test_chat_opens_the_matching_panel():
    sid = _session()
    r = client.post("/api/chat", json={"sid": sid, "message": "I have surface tension vs conc data, want the CMC"}).json()
    assert r["action"]["open"] == "cmc_tensiometry"
    assert "Which do you want" in r["reply"]
    r = client.post("/api/chat", json={"sid": sid, "message": "how bound are the counterions? conductivity data"}).json()
    assert r["action"]["open"] == "cmc_conductivity"
    r = client.post("/api/chat", json={"sid": sid, "message": "hello"}).json()
    assert r["action"] is None


def test_docx_report_follows_the_formatting_spec(full_session):
    r = client.get(f"/api/report/{full_session}.docx?sample=SDS")
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert all(ord(ch) < 128 for ch in text), "keyboard characters only"
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text.strip():
                assert run.font.name == "Times New Roman"
                if run.font.italic:            # the only italics allowed: subsection headings "3.x ..."
                    assert p.text[:2] == "3." and run.font.bold, p.text
    body = [p for p in doc.paragraphs if p.text.startswith("Micelles start forming")]
    assert body and body[0].paragraph_format.first_line_indent is not None
    assert "Short title" in text and "References" in text and "Zana" in text
    footer_xml = doc.sections[0].footer._element.xml
    assert "PAGE" in footer_xml


def test_pdf_report_builds(full_session):
    r = client.get(f"/api/report/{full_session}.pdf?sample=SDS")
    assert r.status_code == 200 and r.content[:5] == b"%PDF-" and len(r.content) > 20000


def test_server_binds_locally_only():
    import inspect
    from surfkit_studio import app as mod
    assert 'host="127.0.0.1"' in inspect.getsource(mod.main)
